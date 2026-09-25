"""
Разбор отчётов площадок (ML Finance → «Отчёты партнёров»).

ЗАЧЕМ ЭТО ВООБЩЕ: у каждой площадки свой файл. Где-то суммы авторских и
смежных лежат отдельными колонками, где-то есть только общая сумма и её надо
поделить, где-то сумма с НДС. В Dista это решалось строкой-формулой вида
`-;-;АРТИКУЛ;-;-;-;КОЛИЧЕСТВО;…` — то есть КОЛОНКИ ПО НОМЕРАМ, с пропусками
через точку с запятой. Стоит площадке переставить столбцы — формула молча
разъезжается, и это ровно то, от чего владелец просил уйти (17.09.2026).

ПОЭТОМУ ЗДЕСЬ КОЛОНКИ ИЩУТСЯ ПО НАЗВАНИЮ. Правило разбора — это «какой
столбец чем является», по именам из шапки, плюс при необходимости формула.
Переставили столбцы местами — разбор не заметит; переименовали — скажет, что
колонки нет, вместо того чтобы взять соседнюю.

ЧТО НА ВЫХОДЕ — ОДИН ФОРМАТ ДЛЯ ВСЕХ: артикул, количество, сумма авторских,
сумма смежных. Дальше по этим четырём числам считается роялти, и расчёту
незачем знать, как выглядел исходный файл (просьба владельца: «свести всё к
одному формату»).

УСТРОЙСТВО ПРАВИЛА (JSON в partner_report_rules.mapping):

    {
      "sku":            {"column": "Артикул"},
      "quantity":       {"column": "Количество"},
      "amount_author":  {"column": "Сумма авт."},
      "amount_related": {"formula": "[Сумма] - [Сумма авт.]"},
      "title":          {"column": "Наименование"}
    }

`column` — взять как есть, `formula` — посчитать по другим колонкам. Формула
намеренно бедная: числа, скобки, + - * / и ссылки на колонки в квадратных
скобках. Это не язык программирования в отчёте, а «сумма минус смежные»;
всё, что сложнее, лучше решать правилом, а не выражением.

НДС снимается ставкой (`vat_rate`), а не формулой в каждой колонке: «сумма с
НДС 20%» — свойство отчёта целиком, и повторять `/1.2` в двух местах значит
однажды поправить одно и забыть второе.
"""
import ast
import codecs
import csv
import io
import re
from datetime import date, datetime, timedelta
from dataclasses import dataclass, field
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from itertools import chain, islice

import openpyxl

from app.countries import country_code

# Поля единого формата. `title` необязателен и нужен только человеку — чтобы в
# предпросмотре было видно, что за трек, если артикул не опознан.
MONEY_FIELDS = ("amount_author", "amount_related")
FIELDS = ("sku", "code", "title", "artist", "quantity", *MONEY_FIELDS)

# ЧЕТЫРЕ ПАРАМЕТРА ОТЧЁТА, КОТОРЫЕ БЫВАЮТ И ПОСТРОЧНЫМИ (24.09.2026, разбор
# настоящего отчёта контрагенту). У МТС и «101 и К» они одни на весь файл и
# задаются значением в правиле. А у Believe в одном отчёте 308 разных
# сочетаний — «Stream / YouTube UGC / streaming / NO», «Creation / TikTok /
# streaming / BY», — и территория идёт по странам. Один снимок на отчёт
# склеил бы их в одну подпись, и в отчёте правообладателю «NO» стало бы
# неотличимо от «MX».
#
# Поэтому те же четыре имени можно указать В ПРАВИЛЕ КАК КОЛОНКУ, и тогда
# значение берётся из строки файла. Не задано колонкой — берётся снимок
# отчёта, как раньше.
ATTR_FIELDS = ("content_type", "usage_type", "usage_kind", "territory")
MAPPABLE_FIELDS = (*FIELDS, *ATTR_FIELDS)
FIELD_LABELS = {
    "sku": "Артикул",
    "code": "Код площадки (ISRC/UPC)",
    "title": "Наименование",
    "artist": "Исполнитель",
    "quantity": "Количество",
    "amount_author": "Сумма авторских",
    "amount_related": "Сумма смежных",
    "content_type": "Тип контента",
    "usage_type": "Тип использования",
    "usage_kind": "Вид использования",
    "territory": "Территория",
}
# Без артикула строку не к чему привязать, без сумм она бессмысленна. Остальное
# необязательно: количество есть не во всех отчётах, название — тем более.
# АРТИКУЛ ЗАДАН ВСЕГДА — колонкой файла ЛИБО «заполняется правилом»
# (`{"auto": True}`). Это не придирка: артикул обязателен, потому что код
# площадки указан верно не всегда, и «пусто» должно означать ровно «ниоткуда
# не берётся» — то есть ошибку настройки, а не молчаливое «как-нибудь
# найдётся» (замечание владельца 24.09.2026).
#
# «Заполняется правилом» значит: в отчёте нашего артикула нет, и трек ищут по
# КОДУ ПЛОЩАДКИ (`code`, отдельное поле — у «101 и К» это «UPC / ISRC»), по
# названию с исполнителем или руками в предпросмотре.
def sku_configured(mapping: dict) -> bool:
    spec = (mapping or {}).get("sku") or {}
    return bool(spec.get("column") or spec.get("auto"))

# Слова, по которым узнаётся ИТОГОВАЯ строка в конце отчёта. У МТС это
# «Итого:», «НДС 22%:», «Итого с НДС:» — строки без кода объекта, но с суммой в
# колонке денег. Не отсечь их значит посчитать выручку дважды и получить
# «строку без артикула» на весь отчёт.
#
# Признак — ВМЕСТЕ: нет артикула И где-то в строке стоит одно из этих слов.
# По одному слову нельзя: «Итого» законно встречается и в названии трека.
# «к выплате» — ВОИС (25.09.2026): под «Итого начислено» у него идёт ещё и
# «К выплате» с той же суммой, и без этого слова отчёт удваивался.
TOTALS_MARKERS = ("итого", "всего", "total", "ндс", "vat", "к выплате")

# Сколько первых строк просматриваем в поисках шапки. У площадок сверху бывает
# шапка-описание на несколько строк (в отчёте МТС, например, данные начинаются
# с восьмой), но не на полсотни.
MAX_HEADER_SCAN = 30

# ТОЧНОСТЬ СТРОКИ — ВОСЕМЬ ЗНАКОВ, а не копейки. Площадки считают дробно (у МТС
# строка «147.0456»), и округление каждой строки до копеек увело итог отчёта на
# 84 копейки от их же «Итого» — на 668 строках набежало. Округляем ОДИН РАЗ, на
# итогах: сверять с платежом человек будет именно их.
#
# Сначала знаков было ЧЕТЫРЕ, и этого не хватило «Зайцев.нет» (24.09.2026):
# цена прослушивания у него 0,024000092 ₽, хвост отбрасывался на каждой строке
# в одну сторону, и на 26 тысячах строк итог ушёл от площадки на 11 копеек.
# Колонки в базе — Numeric(20, 8), миграция d2b87f4c1e69.
PRECISION = Decimal("0.00000001")
CENTS = Decimal("0.01")


# slots=True — НЕ микрооптимизация: этих объектов ровно столько же, сколько
# строк в файле, а в отчёте Believe их 584 тысячи, в отчёте ОМА — полтора
# миллиона. Обычный объект носит с собой словарь атрибутов (~100 байт на
# штуку), и на таком количестве это сотня мегабайт из ничего.
@dataclass(slots=True)
class ReportRow:
    """Строка отчёта, приведённая к единому формату."""

    row_num: int
    sku: str | None = None
    # КОД САМОЙ ПЛОЩАДКИ (ISRC/UPC), если нашего артикула в отчёте нет. По
    # нему трек находится в каталоге, а в `sku` уезжает уже НАШ артикул.
    code: str | None = None
    # Строка без артикула НЕ ошибка разбора: в отчёте МТС таких два десятка —
    # у площадки не проставлен код объекта. Деньги по ним пришли, и молча
    # выкинуть их нельзя; они грузятся, а артикул подбирается по названию и
    # исполнителю (см. find_track_by_name).
    title: str | None = None
    artist: str | None = None
    # Как нашёлся трек: по артикулу из файла или подобран по названию.
    matched_by: str | None = None
    quantity: Decimal | None = None
    amount_author: Decimal = Decimal(0)
    amount_related: Decimal = Decimal(0)
    # Параметры, взятые ИЗ КОЛОНОК файла (см. ATTR_FIELDS). Пусто — значит, у
    # этой площадки они одни на весь отчёт и лежат в его шапке.
    content_type: str | None = None
    usage_type: str | None = None
    usage_kind: str | None = None
    territory: str | None = None
    problems: list = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.problems


@dataclass
class ParseResult:
    """Разобранный файл: шапка, строки, итоги и всё, что пошло не так."""

    columns: list
    header_row: int
    rows: list
    problems: list = field(default_factory=list)
    # ПРЕДУПРЕЖДЕНИЯ — НЕ ОШИБКИ: файл разобран, загрузить его можно, но с ним
    # что-то не так, и человек должен это увидеть до загрузки. `problems`
    # для этого не годятся: они отказывают в загрузке.
    warnings: list = field(default_factory=list)
    # ВАЛЮТЫ СУММ, если правило знает колонку валюты (`mapping["currency"]`).
    # У Believe отчёты приходят в USD, EUR и RUB, и без курса доллары легли бы
    # в базу рублями — ошибка на два порядка, которую на глаз не заметить.
    currencies: set = field(default_factory=set)
    # Отчёт из нескольких файлов: [{name, first_row, last_row}] — какой файл
    # какие (сквозные) номера строк занял. У одного файла пусто.
    files: list = field(default_factory=list)

    @property
    def totals(self) -> dict:
        # Суммы — по ВСЕМ строкам, а не только по беспроблемным: деньги в
        # отчёте есть, даже если у строки нет артикула, и итог должен сходиться
        # с платежом площадки.
        return {
            "rows": len(self.rows),
            "ok_rows": sum(1 for r in self.rows if r.ok),
            "no_sku": sum(1 for r in self.rows if not r.sku),
            "problem_rows": sum(1 for r in self.rows if r.problems),
            "quantity": sum((r.quantity or Decimal(0) for r in self.rows), Decimal(0)),
            "amount_author": rubles(sum((r.amount_author for r in self.rows), Decimal(0))),
            "amount_related": rubles(sum((r.amount_related for r in self.rows), Decimal(0))),
        }


# --------------------------------------------------------------------- чтение


# Сколько байт нюхаем, определяя кодировку. Кириллица в cp1251 ломает проверку
# на utf-8 в первой же строке данных, так что четверти мегабайта хватает с
# огромным запасом, а читать ради этого весь файл — лишний проход по сотне
# мегабайт.
SNIFF_BYTES = 256 * 1024


def detect_encoding(content: bytes) -> str:
    """
    Кодировка текстового отчёта.

    ПРОВЕРЕНО НА НАСТОЯЩИХ ФАЙЛАХ (23.09.2026): выгрузки Believe, Spotify, ОМА
    и «Звука» приходят в cp1251, а не в utf-8. Прежний безусловный
    `decode("utf-8-sig", errors="replace")` превращал в «замену» КАЖДЫЙ
    кириллический символ — на двухстах килобайтах образца до семидесяти тысяч
    штук. Отчёт при этом загружался без единой ошибки и был при этом негоден:
    названия и исполнители приезжали мусором, а это ровно те два поля, по
    которым подбирается артикул, когда площадка не проставила код.

    Порядок проверки важен: cp1251 принимает почти любой байт и потому годится
    только последним. utf-8, наоборот, строг — на кириллице в cp1251 он
    спотыкается сразу.
    """
    if content[:2] in (b"\xff\xfe", b"\xfe\xff"):
        return "utf-16"
    # Инкрементальный декодер не спотыкается о символ, разрезанный границей
    # куска: без final=True он придержит незаконченный хвост, а не сочтёт его
    # ошибкой. Иначе кодировка зависела бы от того, куда попала граница.
    decoder = codecs.getincrementaldecoder("utf-8")()
    try:
        decoder.decode(content[:SNIFF_BYTES], final=False)
        return "utf-8-sig"
    except UnicodeDecodeError:
        return "cp1251"


def _text_stream(content: bytes, encoding: str) -> io.TextIOWrapper:
    """Байты → текстовый поток. Раскодировка идёт по мере чтения, а не разом."""
    return io.TextIOWrapper(
        io.BytesIO(content), encoding=encoding, errors="replace", newline=""
    )


def iter_table(content: bytes, filename: str, sheet: str | None = None):
    """
    Файл → строки таблицы ПО ОДНОЙ, а не списком целиком.

    Excel и текст читаются по-разному, но дальше разбор один: правило не
    должно зависеть от того, прислала площадка .xlsx или .csv.

    ПОЧЕМУ ГЕНЕРАТОР: список держит в памяти весь файл сразу — на отчёте
    Believe это 900 МБ при 137 МБ самого файла, потому что каждая ячейка
    становится отдельным объектом Python. Разбор идёт строка за строкой и
    назад не перематывает, так что копить нечего.
    """
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
        try:
            ws = wb[sheet] if sheet and sheet in wb.sheetnames else wb[wb.sheetnames[0]]
            # РАЗМЕРУ ЛИСТА ИЗ ФАЙЛА НЕ ВЕРИМ (найдено 24.09.2026 на отчёте
            # «Зайцев.нет»). В режиме read_only openpyxl читает ровно столько строк,
            # сколько записано в отметке <dimension> файла, а программа площадки
            # записала туда «8 строк» при двадцати шести тысячах: отчёт разбирался
            # в ноль строк без единой ошибки. reset_dimensions() велит читать до
            # конца данных.
            ws.reset_dimensions()
            for row in ws.iter_rows(values_only=True):
                yield list(row)
        finally:
            # Читателя могут бросить на полпути (см. read_head — ему нужны
            # только верхние строки), и тогда сюда придёт GeneratorExit.
            wb.close()
        return

    encoding = detect_encoding(content)
    # Разделитель угадываем по первой непустой строке: у площадок встречаются
    # и табуляция, и точка с запятой, и запятая. Нюхаем отдельным потоком, а
    # не перематываем рабочий: перемотка текстового потока после обхода
    # итератором — как раз то место, где потом ищут странные ошибки.
    sample = ""
    for line in _text_stream(content, encoding):
        if line.strip():
            sample = line
            break
    delimiter = max(("\t", ";", ","), key=sample.count) if sample else ","
    for row in csv.reader(_text_stream(content, encoding), delimiter=delimiter):
        yield row


def read_table(content: bytes, filename: str, sheet: str | None = None) -> list:
    """
    Весь файл списком.

    Осталась для мелких файлов и тестов; разбор настоящего отчёта идёт через
    `iter_table`, иначе память растёт вместе с файлом.
    """
    return list(iter_table(content, filename, sheet))


def read_head(content: bytes, filename: str, sheet: str | None = None) -> list:
    """
    Верхние строки файла — те, где живут шапка-описание и названия колонок.

    Нужна там, где раньше читался ВЕСЬ файл ради тридцати верхних строк:
    список колонок и период из шапки. На отчёте в полмиллиона строк каждый
    такой проход стоил полторы минуты и гигабайт памяти, а в предпросмотре их
    было три.
    """
    return list(islice(iter_table(content, filename, sheet), MAX_HEADER_SCAN))


def sheet_names(content: bytes, filename: str) -> list:
    """Листы книги — для выбора в интерфейсе. Для текста список пуст."""
    if not (filename or "").lower().endswith((".xlsx", ".xlsm")):
        return []
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    return list(wb.sheetnames)


def _clean(value) -> str:
    return " ".join(str(value if value is not None else "").split())


def normalize_header(value) -> str:
    """
    Ключ сравнения названий колонок: регистр и лишние пробелы значения не
    имеют, как и точки в сокращениях («Сумма авт.» и «Сумма авт»).
    """
    return _clean(value).lower().replace("ё", "е").rstrip(".").strip()


def header_names(row) -> list:
    """
    Названия колонок строки-шапки, где ПОВТОРЫ РАЗВЕДЕНЫ: второй
    «Правообладатель» становится «Правообладатель (2)».

    Зачем. У «101 и К» в шапке две одноимённые колонки: в первой стоит имя
    лейбла («Media Land»), во второй — деньги. Правило адресует колонку по
    названию, и без этого до второй было не дотянуться вовсе — выигрывала
    первая, то есть текст вместо суммы. В списке выбора обе тоже выглядели
    одинаково, и человек не мог понять, какую берёт.

    Суффикс синтетический — в файле его нет, — и это осознанно: он появляется
    только там, где название само по себе перестало быть адресом. Правила,
    написанные раньше, не трогаются: ПЕРВОЕ вхождение остаётся под своим
    именем.
    """
    names, seen = [], {}
    for cell in row:
        name = _clean(cell)
        if not name:
            names.append("")
            continue
        key = normalize_header(name)
        seen[key] = seen.get(key, 0) + 1
        names.append(name if seen[key] == 1 else f"{name} ({seen[key]})")
    return names


def combined_header(row, sub) -> list:
    """
    ШАПКА В ДВЕ СТРОКИ (24.09.2026, отчёт ADV): «Доля Лицензиара, %» сверху
    объединена над «Авторские права» и «Смежные права» снизу, и у второй
    колонки пары верхняя ячейка ПУСТАЯ — по одной верхней строке до неё не
    дотянуться никак. Имя такой колонки — «верх / низ»: «Доля Лицензиара, % /
    Смежные права». Пустой верх берётся от ближайшей колонки слева
    (так выглядит объединённая ячейка), но только если под ним есть подпись;
    колонка без нижней подписи называется, как раньше, одним верхом.

    Числа во второй строке подписью не считаются: у ADV там итоги отчёта.
    """
    top = header_names(row or [])
    subs = [_clean(c) if isinstance(c, str) else "" for c in (sub or [])]
    out, carry = [], ""
    for i in range(max(len(top), len(subs))):
        t = top[i] if i < len(top) else ""
        b = subs[i] if i < len(subs) else ""
        if t:
            carry = t
        base = t or (carry if b else "")
        out.append(f"{base} / {b}" if base and b else (base or b))
    return header_names(out)


def header_at(table: list, i: int, subheader: bool = False) -> list:
    """Имена колонок шапки в строке `i` — с учётом второй строки, если она есть."""
    row = table[i] if i < len(table) else []
    if subheader:
        return combined_header(row, table[i + 1] if i + 1 < len(table) else [])
    return header_names(row)


def guess_header_row(table: list) -> int:
    """
    Где шапка, если правила ещё нет: самая «широкая» текстовая строка сверху.

    Нужна ровно для первого файла нового партнёра — показать человеку список
    колонок, чтобы он собрал правило. Дальше шапку находят по именам из
    правила (find_header_row), и гадать больше не приходится.

    Считаем непустые ячейки и требуем, чтобы это был текст: у шапки их больше
    всего, а строки с данными начинаются ниже и часто содержат числа.
    """
    best_row, best_score = 0, -1
    for i, row in enumerate(table[:MAX_HEADER_SCAN]):
        cells = [c for c in row if _clean(c)]
        if len(cells) < 2:
            continue
        texts = sum(1 for c in cells if not isinstance(c, (int, float)) and parse_number(c) is None)
        score = texts * 10 + len(cells)
        if score > best_score:
            best_row, best_score = i, score
    return best_row


def read_columns(
    content: bytes, filename: str, sheet: str | None = None, head: list | None = None
) -> tuple:
    """
    Названия колонок файла и номер строки с шапкой (0-based).

    Читаем только верх файла: шапка живёт в первых строках, и `guess_header_row`
    дальше MAX_HEADER_SCAN всё равно не смотрит. Готовый `head` можно передать
    снаружи — в предпросмотре он же нужен для поиска периода, и читать его
    дважды незачем.
    """
    head = read_head(content, filename, sheet) if head is None else head
    header_row = guess_header_row(head)
    columns = header_names(head[header_row] if header_row < len(head) else [])
    return columns, header_row


def find_header_row(table: list, wanted: list, subheader: bool = False) -> int:
    """
    Номер строки с шапкой (0-based) — та, где нашлось больше всего нужных
    названий. Ищем по СОДЕРЖИМОМУ, а не по номеру: у площадок сверху бывает
    описание на несколько строк, и число этих строк меняется от файла к файлу.
    """
    wanted_keys = {normalize_header(w) for w in wanted if w}
    best_row, best_hits = 0, -1
    for i, row in enumerate(table[:MAX_HEADER_SCAN]):
        # Имена берём разведёнными: правило может ссылаться на
        # «Правообладатель (2)», и по сырой строке такое не нашлось бы.
        keys = {normalize_header(c) for c in header_at(table, i, subheader) if c}
        hits = len(wanted_keys & keys)
        if hits > best_hits:
            best_row, best_hits = i, hits
        if hits == len(wanted_keys) and wanted_keys:
            break
    return best_row


# ------------------------------------------------------------------- значения


def parse_number(value) -> Decimal | None:
    """
    Число из ячейки: и 1234.56 из Excel, и «1 234,56», и «1,234.56».

    Разделитель решается ПОСЛЕДНИМ знаком: в «1 234,56» это запятая, в
    «1,234.56» — точка. Гадать по стране бессмысленно — отчёты приходят и от
    российских площадок, и от зарубежных.
    """
    if value is None:
        return None
    # ЛОГИЧЕСКАЯ ЯЧЕЙКА — НЕ ЧИСЛО, и это не придирка: в Python `bool` —
    # подкласс `int`, поэтому ИСТИНА из Excel проходила проверку «это число» и
    # роняла `Decimal('True')`, а вместе с ним и разбор всего файла с ошибкой
    # 500. Настоящий случай: отчёт Believe KZ (23.09.2026).
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float, Decimal)):
        # NaN и бесконечность Decimal принимает молча, и дальше они отравили бы
        # итоги: сумма с NaN — это NaN, и объяснить её человеку будет нечем.
        if value != value or value in (float("inf"), float("-inf")):
            return None
        try:
            return Decimal(str(value))
        except InvalidOperation:
            return None
    raw = _clean(value).replace(" ", "").replace(" ", "")
    if not raw:
        return None
    raw = raw.replace("−", "-")          # минус из Word, не ASCII
    last_comma, last_dot = raw.rfind(","), raw.rfind(".")
    if last_comma > last_dot:
        raw = raw.replace(".", "").replace(",", ".")
    else:
        raw = raw.replace(",", "")
    try:
        return Decimal(raw)
    except InvalidOperation:
        return None


_FORMULA_REF = re.compile(r"\[([^\]]+)\]")
_ALLOWED_NODES = (
    ast.Expression, ast.BinOp, ast.UnaryOp, ast.Add, ast.Sub, ast.Mult, ast.Div,
    ast.USub, ast.UAdd, ast.Constant, ast.Name, ast.Load,
)


def formula_columns(expr: str) -> list:
    """Какие колонки упомянуты в формуле — чтобы проверить их наличие заранее."""
    return [m.strip() for m in _FORMULA_REF.findall(expr or "")]


def eval_formula(expr: str, values: dict):
    """
    Значение формулы вида «[Сумма] - [Сумма авт.]».

    Разрешены только числа, скобки и четыре действия: это не язык в отчёте, а
    способ получить сумму, которой в файле нет отдельной колонкой. Всё
    остальное (вызовы, атрибуты, имена) отвергается разбором, а не проверкой
    строки регуляркой — регулярка ловит то, о чём догадался автор, а список
    разрешённых узлов ловит всё остальное по умолчанию.

    None — если хоть одно слагаемое пустое: складывать «нет данных» с числом
    нельзя, такую строку честнее показать проблемной.
    """
    names = {}
    def replace(match):
        key = match.group(1).strip()
        var = "c%d" % len(names)
        names[var] = values.get(normalize_header(key))
        return var

    code = _FORMULA_REF.sub(replace, expr or "")
    if any(v is None for v in names.values()):
        return None
    try:
        tree = ast.parse(code, mode="eval")
    except SyntaxError:
        raise ValueError(f"не разобрать формулу: {expr!r}")
    for node in ast.walk(tree):
        if not isinstance(node, _ALLOWED_NODES):
            raise ValueError(f"в формуле нельзя использовать {type(node).__name__}")
        if isinstance(node, ast.Name) and node.id not in names:
            raise ValueError(f"неизвестное имя в формуле: {node.id}")
        if isinstance(node, ast.Constant) and not isinstance(node.value, (int, float)):
            raise ValueError("в формуле допустимы только числа")
    env = {k: Decimal(str(v)) for k, v in names.items()}
    try:
        return eval(compile(tree, "<formula>", "eval"), {"__builtins__": {}}, env)  # noqa: S307
    except (ZeroDivisionError, InvalidOperation):
        return None


def money(value) -> Decimal:
    """Сумма строки — до восьмого знака (см. PRECISION)."""
    return (value or Decimal(0)).quantize(PRECISION, rounding=ROUND_HALF_UP)


def rubles(value) -> Decimal:
    """Итог — до копейки: столько и переводят."""
    return (value or Decimal(0)).quantize(CENTS, rounding=ROUND_HALF_UP)


# -------------------------------------------------------------------- разбор


def code_text(value) -> str:
    """
    Артикул или код из ячейки — строкой, и целое число без «.0».

    Excel хранит числа дробными: артикул 4100040, набранный числом, читается
    как 4100040.0, и с `tracks.sku` «4100040» он не совпал бы никогда. Так
    приходят номера релизов в отчёте Believe AE (24.09.2026). То же со
    строкой «4100040.0» — её так записывает выгрузка площадки.
    """
    if isinstance(value, bool):
        return _clean(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = _clean(value)
    if _INTEGRAL_FLOAT.fullmatch(text):
        return text.split(".", 1)[0]
    return text


_INTEGRAL_FLOAT = re.compile(r"\d+\.0+")

# ПЕРЕВОДЫ ЗНАЧЕНИЙ ПАРАМЕТРОВ: `{"column": "...", "as": "country"}` в правиле.
# Площадка пишет страну названием, а в отчёте правообладателю нужен код.
TEXT_TRANSFORMS = {"country": country_code}


def currency_factor(rate) -> Decimal:
    """
    Множитель пересчёта валюты в рубли по КУРСУ, как его вписал человек.

    Курс бывает записан ДВУМЯ СПОСОБАМИ, и различаются они величиной (правило
    владельца 24.09.2026): «76,75» — рублей за доллар, на него умножают;
    «0,012216938» — долларов за рубль, на него делят ($13 430,19 / 0,012216938
    = 1 099 309,02 ₽). Курса ровно 1 не бывает, а меньше единицы рублей за
    валюту — тоже: у всех валют, в которых платят площадки, рубль дешевле.
    """
    if rate in (None, ""):
        return Decimal(1)
    value = Decimal(str(rate).replace(",", ".").replace(" ", ""))
    if value <= 0:
        raise ValueError("курс должен быть больше нуля")
    return value if value >= 1 else Decimal(1) / value


def mapping_columns(mapping: dict) -> list:
    """Все названия колонок, которые упоминает правило (прямо или в формуле)."""
    out = []
    for spec in (mapping or {}).values():
        if not isinstance(spec, dict):
            continue
        if spec.get("column"):
            out.append(spec["column"])
        out.extend(formula_columns(spec.get("formula", "")))
    return out


@dataclass(slots=True)
class _Plan:
    """
    Всё, что зависит ТОЛЬКО от правила и шапки, посчитанное один раз.

    Раньше это жило внутри цикла и пересчитывалось на каждой строке: список
    денежных колонок, разбор формул на упомянутые колонки, нормализация
    названий. На отчёте МТС в 657 строк такого не заметить, на Believe в 584
    тысячи — это и есть основное время.

    Вынесено в класс, потому что в одном листе бывает НЕСКОЛЬКО ТАБЛИЦ (у
    Мегафона под основным отчётом идёт отчёт по пакетам со своей шапкой), и
    на её строке план приходится собирать заново.
    """

    fields: list
    sku_col: str | None
    code_col: str | None
    currency_col: str | None
    text_plan: list
    text_only: set
    presence_keys: list
    qty_col: str | None
    qty_formula: str | None
    money_plan: list


def _build_plan(mapping: dict, by_key: dict) -> _Plan:
    """План разбора строки по правилу и шапке (`by_key`: колонка → номер)."""
    wanted = mapping_columns(mapping)
    fields = [
        (key, by_key[key])
        for key in dict.fromkeys(normalize_header(c) for c in wanted)
        if key in by_key
    ]
    sku_spec = (mapping or {}).get("sku") or {}
    code_spec = (mapping or {}).get("code") or {}
    text_plan = [
        (field, normalize_header(spec["column"]), TEXT_TRANSFORMS.get(spec.get("as")))
        for field, spec in (
            (f, (mapping or {}).get(f) or {}) for f in ("title", "artist", *ATTR_FIELDS)
        )
        if spec.get("column")
    ]
    currency_spec = (mapping or {}).get("currency") or {}
    # Колонки ПАРАМЕТРОВ читаются как текст и в `numbers` не попадают: числами
    # они не бывают, а parse_number на каждой строке стоит времени — на отчёте
    # в полмиллиона строк это заметно. Название и исполнителя отсюда
    # исключаем: их колонку теоретически могут упомянуть в денежной формуле,
    # и тогда число понадобится.
    text_only = {
        normalize_header(spec["column"])
        for spec in ((mapping or {}).get(f) or {} for f in (*ATTR_FIELDS, "currency"))
        if spec.get("column")
    }
    qty_spec = (mapping or {}).get("quantity") or {}
    qty_col = normalize_header(qty_spec["column"]) if qty_spec.get("column") else None
    money_plan = []
    for money_key in MONEY_FIELDS:
        spec = (mapping or {}).get(money_key) or {}
        money_plan.append((
            money_key,
            normalize_header(spec["column"]) if spec.get("column") else None,
            spec.get("formula") if not spec.get("column") else None,
            FIELD_LABELS[money_key].lower(),
        ))
    return _Plan(
        fields=fields,
        sku_col=normalize_header(sku_spec["column"]) if sku_spec.get("column") else None,
        code_col=normalize_header(code_spec["column"]) if code_spec.get("column") else None,
        currency_col=(
            normalize_header(currency_spec["column"]) if currency_spec.get("column") else None
        ),
        text_plan=text_plan,
        text_only=text_only,
        # Колонки, по которым решается «в строке есть хоть одно число».
        presence_keys=[
            normalize_header(c)
            for key, spec in (mapping or {}).items()
            if key in (*MONEY_FIELDS, "quantity") and isinstance(spec, dict)
            for c in ([spec["column"]] if spec.get("column") else [])
            + formula_columns(spec.get("formula", ""))
        ],
        qty_col=qty_col,
        qty_formula=qty_spec.get("formula") if not qty_col else None,
        money_plan=money_plan,
    )


def parse_report(
    content: bytes,
    filename: str,
    mapping: dict,
    vat_rate=None,
    sheet: str | None = None,
    limit: int | None = None,
    currency_rate=None,
) -> ParseResult:
    """
    Файл площадки → строки единого формата.

    `vat_rate` — ставка НДС, которую надо ВЫЧЕСТЬ из сумм (20 → делим на 1.2).
    Свойство отчёта целиком, а не отдельной колонки: в одном файле сумма либо
    с налогом, либо без.

    `currency_rate` — курс к рублю для отчёта в валюте (см. `currency_factor`):
    суммы переводятся в рубли ПРИ РАЗБОРЕ, и дальше отчёт ничем не отличается
    от рублёвого — сверка с платежом, выгрузки и расчёт работают как всегда.

    `limit` — сколько строк разобрать (для предпросмотра). Итоги считаются по
    разобранному, и в предпросмотре это честно подписано.
    """
    # ОДИН ПРОХОД ПО ФАЙЛУ: голову придерживаем в памяти (там шапка), хвост
    # дочитываем из того же итератора. Раньше здесь читался весь файл списком,
    # и на отчёте Believe это стоило 900 МБ сверх всего остального.
    rows_iter = iter_table(content, filename, sheet)
    head = list(islice(rows_iter, MAX_HEADER_SCAN))
    wanted = mapping_columns(mapping)
    # Пустое правило — шапку ищем по виду строки, а не по именам: именно так
    # читается первый файл нового партнёра, для которого правила ещё нет.
    # Шапка в две строки (правило с `subheader`) — данные начинаются на
    # строку ниже, а имена колонок собираются из обеих (см. combined_header).
    subheader = bool((mapping or {}).get("subheader"))
    header_row = (
        find_header_row(head, wanted, subheader) if wanted else guess_header_row(head)
    )
    columns = header_at(head, header_row, subheader)
    by_key = {}
    for i, name in enumerate(columns):
        if name:
            by_key.setdefault(normalize_header(name), i)

    result = ParseResult(columns=columns, header_row=header_row, rows=[])

    missing = sorted({c for c in wanted if normalize_header(c) not in by_key})
    if missing:
        result.problems.append(
            "в файле нет колонок: " + ", ".join(f"«{c}»" for c in missing)
        )
        return result
    if not sku_configured(mapping):
        result.problems.append(
            "в правиле не задан «%s»: выберите колонку или «заполняется правилом»"
            % FIELD_LABELS["sku"]
        )
        return result

    divisor = Decimal(1)
    if vat_rate:
        divisor = Decimal(1) + Decimal(str(vat_rate)) / Decimal(100)
    # Курс — отдельным множителем, а НДС по-прежнему делением: свернуть их в
    # одно число нельзя, 1/1.22 — бесконечная дробь, и на границе округления
    # строка изредка расходилась бы с прежним расчётом.
    try:
        factor = currency_factor(currency_rate)
    except (ValueError, InvalidOperation):
        result.problems.append(f"курс «{currency_rate}» — это не число больше нуля")
        return result

    data = chain(head[header_row + 1 + subheader:], rows_iter)
    if limit is not None:
        data = islice(data, limit)

    plan = _build_plan(mapping, by_key)

    # ВТОРАЯ ТАБЛИЦА В ТОМ ЖЕ ЛИСТЕ. У Мегафона под основным отчётом идёт
    # отчёт по пакетам: своя шапка, свои колонки, свои формулы. Строки его
    # данных под правилом первой таблицы читались бы как мусор — с чужими
    # числами в денежных колонках.
    #
    # Узнаём по СОДЕРЖИМОМУ строки: встретили строку, где есть все колонки
    # дополнительного правила, — дальше разбираем по нему. Тот же приём, что
    # и с поиском шапки: на номера строк полагаться нельзя, длина первой
    # таблицы меняется от месяца к месяцу.
    tables = [
        (t, {normalize_header(c) for c in mapping_columns(t)})
        for t in (mapping or {}).get("tables") or []
    ]

    for offset, raw in enumerate(data):
        row_num = header_row + 2 + subheader + offset   # как в Excel: с единицы, с шапкой
        if not any(c is not None and str(c).strip() for c in raw):
            continue

        if tables:
            keys = {normalize_header(c) for c in header_names(raw) if c}
            extra = next((t for t, wanted_keys in tables if wanted_keys <= keys), None)
            if extra is not None:
                here = {}
                for i, name in enumerate(header_names(raw)):
                    if name:
                        here.setdefault(normalize_header(name), i)
                plan = _build_plan(extra, here)
                continue

        size = len(raw)
        values = {key: (raw[i] if i < size else None) for key, i in plan.fields}

        row = ReportRow(row_num=row_num)
        row.sku = (code_text(values.get(plan.sku_col)) or None) if plan.sku_col else None
        row.code = (code_text(values.get(plan.code_col)) or None) if plan.code_col else None
        for text_field, column, transform in plan.text_plan:
            text = _clean(values.get(column)) or None
            setattr(row, text_field, transform(text) if transform and text else text)
        if plan.currency_col:
            currency = _clean(values.get(plan.currency_col)).upper()
            if currency:
                result.currencies.add(currency)

        numbers = {
            key: parse_number(value)
            for key, value in values.items()
            if key not in plan.text_only
        }

        # СТРОКА БЕЗ ЕДИНОГО ЧИСЛА — НЕ ДАННЫЕ. В конце отчёта МТС идёт блок
        # подписи: «ОТ ЛИЦЕНЗИАРА», «_______ /_______/», «М.П.» — они попадают
        # в таблицу как строки с мусором вместо артикула. Пустая ячейка и ноль
        # здесь разные вещи: ноль — это данные (площадка честно сообщает, что
        # денег не было), пустота — оформление.
        if not any(numbers.get(key) is not None for key in plan.presence_keys):
            continue

        if plan.qty_col:
            row.quantity = numbers.get(plan.qty_col)
        elif plan.qty_formula:
            row.quantity = _apply_formula(plan.qty_formula, numbers, row, "количество")

        for money_key, column, formula, label in plan.money_plan:
            value = None
            if column:
                value = numbers.get(column)
                if value is None and _clean(values.get(column)):
                    row.problems.append(
                        f"{label}: «{_clean(values.get(column))}» — это не число"
                    )
            elif formula:
                value = _apply_formula(formula, numbers, row, label)
            setattr(row, money_key, money((value or Decimal(0)) * factor / divisor))

        # Итоговая строка в конце файла — не данные: пропускаем целиком, иначе
        # её сумма удвоит отчёт.
        if not row.sku and not row.code:
            text = " ".join(_clean(c).lower() for c in raw)
            if any(marker in text for marker in TOTALS_MARKERS):
                continue
        result.rows.append(row)

    _warn_if_lossy(result)
    return result


# «?» ВНУТРИ СЛОВА — след потерянной буквы, а не знак вопроса. Настоящий
# вопросительный знак стоит в конце («Кто?», «Где ты?»), а между двумя буквами
# он оказывается тогда, когда файл сохранили в кодировку, которая эту букву не
# вмещает: так пропадают казахские ә, ғ, қ, ң, ө, ұ, ү, і, турецкие ı и ğ,
# скандинавские ø и å, немецкое ä.
LOST_LETTER = re.compile(r"\w\?\w", re.UNICODE)
# Порог: и по числу строк, и по доле. Одна-две такие строки бывают и в
# честном файле (название вида «Что?Где?Когда»), а вот сотые доли процента на
# большом отчёте — это уже потеря.
LOSSY_MIN_ROWS = 10
LOSSY_MIN_SHARE = 0.002


def _warn_if_lossy(result: ParseResult) -> None:
    """
    Сказать человеку, если файл приехал с уже потерянными буквами.

    ПРОВЕРЕНО НА НАСТОЯЩИХ ФАЙЛАХ (23.09.2026): у одного и того же отчёта
    версия .txt и версия .xlsx различаются именно этим. В Spotify за январь
    текстовый вариант содержит 21 015 строк с «?» вместо букв, а книга Excel —
    две, и в ней целы Ğ, ş, ö, Ø. То же у Believe KZ (8 357 против нуля) и у
    ОМА. Текстовый вариант сохранён в cp1251, а в ней этих букв просто нет.

    Починить такой файл нечем: «?» не помнит, какая буква там была. Поэтому
    единственное полезное действие — предупредить ДО загрузки, пока человек
    может взять другой вариант того же отчёта.
    """
    if not result.rows:
        return
    lost = sum(
        1 for row in result.rows
        if (row.title and LOST_LETTER.search(row.title))
        or (row.artist and LOST_LETTER.search(row.artist))
    )
    if lost < LOSSY_MIN_ROWS or lost < len(result.rows) * LOSSY_MIN_SHARE:
        return
    result.warnings.append(
        f"В файле {lost} строк, где буквы заменены на «?» — он сохранён в "
        "кодировке, которая их не вмещает (так получается при сохранении в "
        "«Текст с разделителями»). Загрузить можно, но названия и исполнители "
        "приедут испорченными. Если у этого отчёта есть вариант .xlsx, "
        "возьмите его: в нём буквы обычно целы."
    )


def _apply_formula(expr: str, numbers: dict, row: ReportRow, label: str):
    try:
        value = eval_formula(expr, numbers)
    except ValueError as exc:
        row.problems.append(f"{label}: {exc}")
        return None
    if value is None:
        row.problems.append(f"{label}: в формуле пустое значение")
    return value


# ------------------------------------------------- подбор трека по названию

# Служебные слова в поле исполнителя: они не различают артистов, а только
# связывают их («ОСОБОВ feat. TRUEтень»). Сравнивать по ним нельзя, иначе
# «Slim & Константа» и «Slim, Константа» окажутся разными, а это один дуэт.
ARTIST_STOPWORDS = {"feat", "ft", "featuring", "prod", "vs", "and", "x"}


def normalize_name(value) -> str:
    """
    Ключ сравнения названий: регистр, «ё», знаки препинания и лишние пробелы
    значения не имеют. «Тмстс!» и «тмстс» — одно и то же название.
    """
    text = str(value or "").lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я]+", " ", text)
    return " ".join(text.split())


def artist_tokens(value) -> list:
    """
    Исполнитель → набор значимых слов. Разделители («&», «,», «feat.») и
    регистр отброшены: в отчёте пишут «Slim & Константа», в каталоге —
    «Slim, Константа», и это один и тот же дуэт.
    """
    words = normalize_name(value).split()
    return sorted(w for w in words if w not in ARTIST_STOPWORDS)


def artists_match(left, right) -> bool:
    """
    Совпадают ли исполнители — СТРОГО, но с поправкой на то, как их пишут.

    Разрешены: другой регистр, другой разделитель, сокращение имени до
    инициала («Ю. Шатунов» против «Юрий Шатунов»). Всё остальное — разные
    артисты: «SLIMUS, Константа» и «Slim, Константа» на одном названии
    «Азимут» — это два РАЗНЫХ трека в каталоге, и подставить наугад любой из
    них значит отправить чужие деньги.

    Поэтому наборы слов должны совпасть ПОЛНОСТЬЮ (по одному слову на слово),
    а не «одно входит в другое»: «ОСОБОВ» и «ОСОБОВ feat. TRUEтень» —
    разные исполнители, хоть первый и содержится во втором.
    """
    a, b = artist_tokens(left), artist_tokens(right)
    if not a or not b or len(a) != len(b):
        return False
    rest = list(b)
    for word in a:
        pair = None
        for candidate in rest:
            if candidate == word:
                pair = candidate
                break
            # инициал против полного слова: «ю» и «юрий»
            if len(word) == 1 and candidate.startswith(word):
                pair = candidate
                break
            if len(candidate) == 1 and word.startswith(candidate):
                pair = candidate
                break
        if pair is None:
            return False
        rest.remove(pair)
    return True


def search_word(title) -> str:
    """
    Самое длинное слово названия — по нему ищем кандидатов в каталоге.

    Искать по всему названию нельзя: в отчёте оно бывает со знаком вопроса, в
    каталоге без, — и точное сравнение промахнётся. Слово даёт короткий список
    кандидатов, а решает уже строгое сравнение целиком.
    """
    words = [w for w in normalize_name(title).split() if len(w) >= 3]
    return max(words, key=len, default="")


def pick_track(title, artist, candidates) -> object | None:
    """
    Единственный сильный кандидат — или ничего.

    `candidates` — записи каталога (объекты с title и artist). Совпасть должны
    ОБА поля; если подошли несколько разных треков, не выбираем ни одного:
    угаданный артикул хуже пустого, потому что деньги уедут молча и не туда.
    """
    key = normalize_name(title)
    if not key or not str(artist or "").strip():
        return None
    hits = [
        c
        for c in candidates
        if normalize_name(c.title) == key and artists_match(artist, c.artist)
    ]
    unique = {c.id for c in hits}
    return hits[0] if len(unique) == 1 else None


MONTHS_RU = (
    "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
    "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь",
)
QUARTERS_RU = ("I", "II", "III", "IV")


# ПРИПИСКА ТЕРРИТОРИИ К ПЛОЩАДКЕ (просьба владельца 25.09.2026). Believe
# шлёт три отчёта за месяц — RU, KZ и AE, — а в справочнике это одна
# площадка, и после привязки к поступлению (суммы уже в рублях) три строки
# списка стали неразличимы. Различает их ВАЛЮТА ИСХОДНОГО ОТЧЁТА: RU — в
# рублях, KZ — в евро, AE — в долларах (сверено по всем девяти загруженным).
# Валюта хранится снимком и после привязки не меняется, поэтому приписка
# держится. Ключ — имя площадки без учёта регистра.
REPORT_REGIONS = {
    "beleive digital": {"RUB": "RU", "EUR": "KZ", "USD": "AE"},
}


def report_region(partner_name: str | None, currency: str | None) -> str | None:
    """«RU» / «KZ» / «AE» для площадок с несколькими отчётами за период."""
    regions = REPORT_REGIONS.get((partner_name or "").strip().casefold())
    return regions.get(currency or "RUB") if regions else None


def period_label(start, end) -> str:
    """
    Подпись периода отчёта: «Июль 2026», «III кв. 2026» или «01.06.2026 —
    15.07.2026».

    Собирает СЕРВЕР, как и подпись периода у поступлений: формат — правило, и
    разъезжаться ему между экраном, списком и будущей выгрузкой незачем.
    Месяц и квартал узнаются по датам, а не хранятся отдельным признаком:
    признак пришлось бы поддерживать в согласии с датами, а даты и так всё
    говорят.
    """
    if start is None or end is None:
        return ""
    import calendar

    last_day = calendar.monthrange(end.year, end.month)[1]
    whole_months = start.day == 1 and end.day == last_day and start.year == end.year
    if whole_months and start.month == end.month:
        return f"{MONTHS_RU[start.month - 1]} {start.year}"
    if whole_months and (start.month - 1) % 3 == 0 and end.month == start.month + 2:
        return f"{QUARTERS_RU[(start.month - 1) // 3]} кв. {start.year}"
    return f"{start:%d.%m.%Y} — {end:%d.%m.%Y}"


# ------------------------------------------------- готовые правила площадок

# Колонки отчёта МТС. Выписаны полностью и нарочно: правило должно совпадать
# с файлом буква в букву, а «примерно такая колонка» — это ровно тот случай,
# когда деньги молча считаются не из того столбца.
_MTS_SKU = "Код объекта Контента"
_MTS_AMOUNT = "Сумма вознаграждения Лицензиара, руб (без НДС)"
_MTS_RATE_AUTHOR = "Ставка вознаграждения Лицензиара за авторские права на Произведение,%"
_MTS_RATE_RELATED = (
    "Ставка вознаграждения Лицензиара за смежные права на Исполнение и Фонограмму,%"
)

# МЕГАФОН. Имена колонок длинные, и в формулах они повторяются по четыре
# раза — держим их здесь, чтобы опечатка не разъехалась между авторскими и
# смежными.
# Колонки отчёта ADV (двойные пробелы внутри имён — как в файле; сравнение
# идёт без учёта лишних пробелов).
_ADV_CODE = "Код Произведения"
_ADV_KIND = "Вид Контента"
_ADV_QTY = "Кол-во Загрузок/ Прослушиваний (штук)"
_ADV_PRICE = "Стоимость Контента (рубли, без учета НДС)"
_ADV_SHARE = "Доля Лицензиара, %"
_ADV_RATE = "Вознаграждение Лицензиара (%/ руб.)"
_MF_PRICE = "Стоимость загрузки, руб. без НДС"
_MF_RATE_AUTHOR = "Ставка Лицензиара (авторские)"
_MF_RATE_RELATED = "Ставка Лицензиара (смежные)"
_MF_SHARE_AUTHOR = "Доля авторских прав, %"
_MF_SHARE_RELATED = "Доля смежных прав, %"
# Вторая таблица того же листа — отчёт по пакетам.
_MF_PACK_PRICE = "Стоимость Запроса Пакета, руб. без НДС"
_MF_PACK_RATE_AUTHOR = "Вознаграждение Лицензиара (авторские права), %"
_MF_PACK_RATE_RELATED = "Вознаграждение Лицензиара (смежные права), %"
_MF_PACK_OURS = "Кол-во единиц Мобильного контента Правообладателя в пакете"
_MF_PACK_TOTAL = "Общее кол-во единиц Мобильного контента в Пакете"

# Колонки отчёта «Зайцев.нет»: названия длинные, и в правиле их удобнее
# держать константами, чтобы приметы и соответствие не разошлись опечаткой.
_ZN_SKU = "Код Контента Лицензиара"
_ZN_SHOP = "Витрина, где использовался Контент"
_ZN_QUANTITY = "Кол-во Загрузок/Прослушиваний"
_ZN_AUTHOR = (
    "Сумма лицензионного вознаграждения Лицензиара за авторские права"
    "(руб., без учета НДС)"
)
_ZN_RELATED = (
    "Сумма лицензионного вознаграждения Лицензиара за смежные права"
    "(руб., без учета НДС)"
)

BUILTIN_RULES = (
    {
        "name": "МТС",
        # Приметы файла: по ним правило и узнаётся. Не по имени партнёра —
        # площадку могут завести и как «МТС», и как «MTS Медиа», а колонки
        # у её отчёта всегда одни и те же.
        "signature": (_MTS_SKU, _MTS_AMOUNT, _MTS_RATE_AUTHOR, _MTS_RATE_RELATED),
        # Кому в справочнике партнёров принадлежит такой отчёт. Имя ТОЧНОЕ, без
        # «начинается на»: в справочнике живут ещё «МТС Авторские» и «МТС
        # Беларусь», и подставить не ту площадку — это деньги не на тот счёт.
        # Подсказка так и остаётся подсказкой: партнёра видно в поле и его
        # можно сменить.
        "partner_names": ("МТС",),
        # НДС в отчёте НЕ заложен: все денежные колонки подписаны «без НДС»,
        # а налог площадка добавляет итоговой строкой в конце файла.
        "vat_rate": None,
        # Параметры отчёта МТС — как они стоят в Dista (скриншот владельца
        # 18.09.2026). Это ЗАГОТОВКА: подставляется, пока у партнёра не
        # сохранено своё, и правится прямо в форме.
        "attributes": {
            "content_type": "RBT",
            # Прочерк, а не «<не участвует>»: в отчёте это графа, которую
            # площадка не заполняет, и прочерк читается так же, а выглядит
            # как значение, а не как служебная пометка.
            "usage_type": "—",
            "usage_kind": "Mobile",
            "territory": "RU",
        },
        "mapping": {
            "sku": {"column": _MTS_SKU},
            "title": {"column": "Название объекта Контента"},
            "artist": {"column": "Исполнитель"},
            "quantity": {"column": "Кол-во Продаж"},
            # Отдельных колонок с авторскими и смежными у МТС нет: есть общая
            # сумма вознаграждения и две ставки, по которым она делится.
            # Сверено с колонками «авт»/«смж», посчитанными владельцем вручную:
            # 664 строки, расхождений ноль.
            "amount_author": {
                "formula": f"[{_MTS_AMOUNT}] * [{_MTS_RATE_AUTHOR}]"
                           f" / ([{_MTS_RATE_AUTHOR}] + [{_MTS_RATE_RELATED}])"
            },
            "amount_related": {
                "formula": f"[{_MTS_AMOUNT}] * [{_MTS_RATE_RELATED}]"
                           f" / ([{_MTS_RATE_AUTHOR}] + [{_MTS_RATE_RELATED}])"
            },
        },
    },
    {
        "name": "101 и К",
        # Приметы: «Вал» и «Переработчик» вместе не встречаются больше нигде.
        # «Правообладатель» в приметы не берём — этим словом в шапке названы
        # ДВЕ колонки, и как примета оно ничего не различает.
        "signature": ("Имя", "Название", "UPC / ISRC", "Вал", "Переработчик"),
        "partner_names": ("101 и К",),
        # СУММЫ В ОТЧЁТЕ С НДС 22%: сверено с готовым файлом владельца —
        # 2436.4736 / 1.22 = 1997.1095, и так все четыре строки.
        "vat_rate": 22,
        # Параметры — со скриншота Dista (владелец 24.09.2026). Прочерк вместо
        # «<не участвует>»: это графа, которую площадка не заполняет.
        "attributes": {
            "content_type": "—",
            "usage_type": "—",
            "usage_kind": "streaming",
            "territory": "CIS/RU",
        },
        "mapping": {
            # НАШЕГО АРТИКУЛА В ОТЧЁТЕ НЕТ — есть «UPC / ISRC» одной ячейкой
            # («3617380567893 / DG-A0P-23-16666»). Поэтому колонка ложится в
            # КОД ПЛОЩАДКИ, а не в артикул: по коду трек находится в каталоге
            # (_code_candidates в роутере), и в строку уезжает НАШ артикул.
            # У строки без кода остаётся подбор по названию и исполнителю —
            # так находится «Ёлка — Понедельник», у которой ячейка пуста.
            # Нашего артикула в отчёте НЕТ — он находится по коду площадки
            # и подставляется в строку. Признак явный: «пусто» у артикула
            # означает ошибку настройки, а не «как-нибудь найдётся».
            "sku": {"auto": True},
            "code": {"column": "UPC / ISRC"},
            "title": {"column": "Название"},
            "artist": {"column": "Имя"},
            # КОЛИЧЕСТВО ВСЕГДА 1 (правило владельца): отчёт сводный, в нём
            # строка на трек за квартал, а не продажи поштучно.
            "quantity": {"formula": "1"},
            # ДЕНЬГИ — КОЛОНКА «Правообладатель», ВТОРАЯ ПО СЧЁТУ. Первая с
            # тем же названием держит имя лейбла («Media Land»), и без
            # разведения повторов (header_names) правило читало бы текст
            # вместо суммы.
            #
            # Всё уходит в СМЕЖНЫЕ (решение владельца 24.09.2026): эта доля —
            # доля владельца фонограммы, а авторская у DFM, она стоит в
            # соседней колонке «Паблишер». Авторских поэтому нет вовсе —
            # отсутствующее поле сервер считает нулём.
            "amount_related": {"column": "Правообладатель (2)"},
        },
    },
    {
        # ADV (АдвМьюзик), 25.09.2026, образец владельца «ADV май 26.xlsx»:
        # 4 049 строк, «Всего» 340 302,73 без НДС — сходится до копейки.
        "name": "ADV",
        "signature": (_ADV_CODE, _ADV_KIND, _ADV_QTY, _ADV_PRICE),
        "partner_names": ("ADV Music",),
        # Суммы в отчёте БЕЗ НДС — так подписана колонка; НДС идёт отдельной
        # строкой под итогом.
        "vat_rate": None,
        # Параметры — от владельца (25.09.2026, как в Dista): тип контента —
        # колонка отчёта (в каждой строке свой, «MP3 фонограмма» или «РБТ»,
        # поэтому он в mapping), тип использования не участвует (прочерк, как
        # у МегаФона), вид — streaming, территория — CIS;RU, как у Зайцев.нет.
        "attributes": {
            "usage_type": "—",
            "usage_kind": "streaming",
            "territory": "CIS;RU",
        },
        "mapping": {
            # ШАПКА В ДВЕ СТРОКИ: доли и ставки лицензиара разбиты под общей
            # подписью на авторские и смежные (см. combined_header). Признак
            # ВНУТРИ mapping, как `tables` и `period`: так он едет вместе с
            # правилом и в сохранённое у партнёра.
            "subheader": True,
            "sku": {"column": _ADV_CODE},
            "title": {"column": "Название"},
            "artist": {"column": "Исполнитель"},
            "quantity": {"column": _ADV_QTY},
            "content_type": {"column": _ADV_KIND},
            # ГОТОВЫХ СУММ ПО ВИДАМ ПРАВ НЕТ — есть только общая «Сумма
            # вознаграждения». Её площадка считает как стоимость × количество
            # × (доля авторских × ставка авторских + доля смежных × ставка
            # смежных), и каждое слагаемое — ровно наша сумма по виду права.
            # Сверено на всех 4 049 строках образца: расхождений ноль.
            "amount_author": {
                "formula": f"[{_ADV_PRICE}] * [{_ADV_QTY}]"
                           f" * [{_ADV_SHARE} / Авторские права]"
                           f" * [{_ADV_RATE} / за авторские права]"
            },
            "amount_related": {
                "formula": f"[{_ADV_PRICE}] * [{_ADV_QTY}]"
                           f" * [{_ADV_SHARE} / Смежные права]"
                           f" * [{_ADV_RATE} / за смежные права]"
            },
        },
    },
    {
        # ВОИС (25.09.2026, образцы владельца: отчёты по договорам 1766/12
        # «изготовителю фонограмм» и 3197/14 «исполнителю» — формат один).
        # НАШЕГО АРТИКУЛА В ОТЧЁТЕ НЕТ: раньше владелец проставлял его руками.
        # Теперь — по ISRC, а где его нет (больше половины строк), по названию
        # и исполнителю: обычный порядок привязки, `sku` → «заполняется
        # правилом». Что не нашлось, впишут в предпросмотре — и площадка это
        # запомнит.
        "name": "ВОИС",
        "signature": ("Наименование фонограммы", "Изготовитель фонограммы", "ISRC",
                      "Сумма вознаграждения руб."),
        "partner_names": ("ВОИС",),
        # НДС ВНУТРИ СУММ И ЕГО НАДО СНЯТЬ (владелец): «Итого начислено
        # 2 737,39, в т.ч. НДС 493,63» — это 22%. Так же делил и образец для
        # Dista: 0,13 / 1,22 = 0,1066.
        "vat_rate": 22,
        # Параметры — со скриншота Dista владельца.
        "attributes": {
            "content_type": "—",
            "usage_type": "—",
            "usage_kind": "Public performance",
            "territory": "RU",
        },
        "mapping": {
            "sku": {"auto": True},
            "code": {"column": "ISRC"},
            "title": {"column": "Наименование фонограммы"},
            "artist": {"column": "Исполнитель"},
            # Прослушиваний ВОИС не сообщает — по строке на трек, как в
            # образце для Dista («количество 1»).
            "quantity": {"formula": "1"},
            # Деньги за использование ФОНОГРАММЫ — смежные права.
            "amount_related": {"column": "Сумма вознаграждения руб."},
        },
    },
    {
        # ВОИС В СТАРОЙ ФОРМЕ — таблица, которую владелец собирал руками для
        # загрузки в Dista («4. ВОИС_образец.xlsx»: Код, Количество, Сумма;
        # формула Dista «АРТИКУЛ;КОЛИЧЕСТВО;СУММА_СМЖ»). Артикулы в ней уже
        # проставлены, а суммы уже БЕЗ НДС (разделены на 1,22) — делить
        # второй раз нельзя.
        "name": "ВОИС (таблица для Dista)",
        "signature": ("Код", "Количество", "Сумма"),
        # Три слова, которые найдутся в любой самодельной таблице: узнаём
        # файл, только если в шапке ровно они (см. `exact` в match_builtin).
        "exact": True,
        "partner_names": ("ВОИС",),
        "vat_rate": None,
        "attributes": {
            "content_type": "—",
            "usage_type": "—",
            "usage_kind": "Public performance",
            "territory": "RU",
        },
        "mapping": {
            "sku": {"column": "Код"},
            "quantity": {"column": "Количество"},
            "amount_related": {"column": "Сумма"},
        },
    },
    {
        "name": "МегаФон",
        # Приметы — колонки ОСНОВНОЙ таблицы: «Тип Мобильного контента» и обе
        # ставки лицензиара вместе не встречаются больше нигде.
        "signature": (_MF_PRICE, _MF_RATE_AUTHOR, _MF_RATE_RELATED, "Кол-во загрузок"),
        "partner_names": ("МегаФон",),
        # Суммы в отчёте БЕЗ НДС: колонка так и подписана, а строка «Итого с
        # НДС» идёт отдельно в самом низу.
        "vat_rate": None,
        "attributes": {
            "content_type": "RBT",
            "usage_type": "—",
            "usage_kind": "Mobile",
            "territory": "RU",
        },
        "mapping": {
            "sku": {"column": "Код"},
            "title": {"column": "Наименование"},
            "artist": {"column": "Исполнитель"},
            "quantity": {"column": "Кол-во загрузок"},
            # ОТДЕЛЬНЫХ КОЛОНОК С СУММАМИ НЕТ: есть общая «Сумма
            # вознаграждения», а авторские и смежные считаются из цены,
            # ставки, количества и ДОЛИ. Доля обязательна: в майском отчёте
            # три строки из 206 идут с долей 0.4, 0.25 и 0.5 при нулевой доле
            # смежных, и без неё суммы по ним завышались бы вдвое-вчетверо.
            # Сверено с расчётом владельца: 206 строк, расхождений ноль.
            "amount_author": {
                "formula": f"[{_MF_PRICE}] * [{_MF_RATE_AUTHOR}]"
                           f" * [Кол-во загрузок] * [{_MF_SHARE_AUTHOR}]"
            },
            "amount_related": {
                "formula": f"[{_MF_PRICE}] * [{_MF_RATE_RELATED}]"
                           f" * [Кол-во загрузок] * [{_MF_SHARE_RELATED}]"
            },
            # ВТОРАЯ ТАБЛИЦА ТОГО ЖЕ ЛИСТА — «Отчет распространения пакетов».
            # У неё своя шапка ниже основной, свои колонки и своя формула, а в
            # готовом отчёте владельца её не было вовсе: считали только
            # основную. Теперь учитываем обе (просьба владельца 24.09.2026).
            #
            # ФОРМУЛА ВЫВЕДЕНА ИЗ САМОГО ФАЙЛА, а не придумана: у площадки
            # есть своя колонка «Сумма Вознаграждения Лицензиара, руб. без
            # НДС», и произведение ниже сходится с ней на всех трёх строках
            # (8.64, 40.67, 3.50). Деньги за пакет делятся по числу вещей в
            # нём — отсюда доля «наших единиц» к общему числу.
            "tables": [
                {
                    "sku": {"column": "Код"},
                    "title": {"column": "Наименование Произведения/Фонограммы в составе Пакета"},
                    "artist": {"column": "Исполнитель"},
                    "quantity": {"column": "Кол-во Запросов"},
                    "amount_author": {
                        "formula": f"[{_MF_PACK_PRICE}] * [Кол-во Запросов]"
                                   f" * [{_MF_PACK_OURS}] / [{_MF_PACK_TOTAL}]"
                                   f" * [{_MF_PACK_RATE_AUTHOR}] * [{_MF_SHARE_AUTHOR}]"
                    },
                    "amount_related": {
                        "formula": f"[{_MF_PACK_PRICE}] * [Кол-во Запросов]"
                                   f" * [{_MF_PACK_OURS}] / [{_MF_PACK_TOTAL}]"
                                   f" * [{_MF_PACK_RATE_RELATED}] * [{_MF_SHARE_RELATED}]"
                    },
                }
            ],
        },
    },
    {
        "name": "Зайцев.нет",
        # Приметы: «Витрина» и «Код Контента Лицензиара» вместе не
        # встречаются больше ни у кого (образец владельца 24.09.2026, отчёт
        # за май 2026, 26 383 строки).
        "signature": (_ZN_SKU, _ZN_SHOP, _ZN_QUANTITY, _ZN_AUTHOR, _ZN_RELATED),
        # В справочнике площадка записана заглавными; сравнение идёт casefold.
        "partner_names": ("ЗАЙЦЕВ.НЕТ",),
        # Суммы БЕЗ НДС: так подписаны все денежные колонки, а «Сумма НДС
        # 22%» и «в том числе НДС» идут отдельными строками в самом низу.
        "vat_rate": None,
        # Параметры — со скриншота Dista (владелец 24.09.2026). Прочерк вместо
        # «<не участвует>», как у остальных площадок. Территория записана как
        # в Dista, через точку с запятой.
        "attributes": {
            "content_type": "—",
            "usage_type": "—",
            "usage_kind": "streaming",
            "territory": "CIS;RU",
        },
        "mapping": {
            # НАШ АРТИКУЛ — «Код Контента Лицензиара»: площадка его знает. Но
            # у 1 116 строк из 26 383 он пуст, и тогда трек ищется по ISRC
            # (колонка кода площадки), а не нашёлся — по названию и
            # исполнителю. Колонка UPC в файле почти всегда пуста и в поиске
            # не нужна: ISRC точнее.
            "sku": {"column": _ZN_SKU},
            "code": {"column": "ISRC код"},
            "title": {"column": "Название Контента"},
            "artist": {"column": "Исполнитель"},
            "quantity": {"column": _ZN_QUANTITY},
            # ОБЕ СУММЫ ЕСТЬ ГОТОВЫМИ КОЛОНКАМИ — формула не нужна. Сверено с
            # итогом самого файла: авторские 7 909,11, смежные 39 545,55, всего
            # 47 454,67, до копейки.
            "amount_author": {"column": _ZN_AUTHOR},
            "amount_related": {"column": _ZN_RELATED},
        },
    },
    # BELIEVE (24.09.2026, образцы владельца — отчёты за май 2026: RU на 596
    # тыс. строк, KZ на 280 тыс., AE на 86 тыс.). Правил ДВА, потому что шапка
    # у отчётов разная: RU подписан по-английски, KZ и AE — по-русски, а
    # столбцы и смысл у них одни и те же. Сверено с итогами файлов до копейки:
    # RU 4 858 741,81 ₽, KZ 10 456,59 €, AE 2 950,91 $.
    {
        "name": 'Believe',
        "signature": ('Release Catalog nb', 'Client Payment Currency', 'Net Revenue', 'Sales Type', 'Platform'),
        # Площадка в справочнике одна на все три отчёта (RU, KZ, AE): платежи
        # от Believe приходят от «BELEIVE DIGITAL» (опечатка в справочнике
        # живёт давно, и на неё завязано сопоставление платежей).
        "partner_names": ("BELEIVE DIGITAL",),
        "vat_rate": None,
        # Вид использования у Believe один на все строки — «streaming», как в
        # детализированном отчёте правообладателю из Dista. Остальные три
        # параметра идут ИЗ КОЛОНОК: у Believe в одном отчёте сотни сочетаний
        # площадки, типа продажи и страны.
        "attributes": {
            "usage_kind": "streaming",
        },
        "mapping": {
            # АРТИКУЛ — КАТАЛОЖНЫЙ НОМЕР РЕЛИЗА, и он главный (решение
            # владельца 24.09.2026: «именно по нашему артикулу мы сверяем в
            # первую очередь»). У альбома, сборника и YouTube-канала это номер
            # всей позиции (4100008 — канал Александра Иванова, 200 с лишним
            # видео), и деньги ложатся на неё, как в Dista. ISRC — запасной
            # путь, когда номера релиза нет в каталоге или он пуст.
            "sku": {"column": 'Release Catalog nb'},
            "code": {"column": 'ISRC'},
            "title": {"column": 'Track title'},
            "artist": {"column": 'Artist Name'},
            "quantity": {"column": 'Quantity'},
            # ВСЁ — В СМЕЖНЫЕ: «Сумма вознаграждения» (Net Revenue) по формуле
            # Dista владельца. Механика (Mechanical Fee) не берётся: она почти
            # всегда ноль (39 строк на 534 ₽ в RU за май).
            "amount_related": {"column": 'Net Revenue'},
            # Параметры — по формуле Dista: площадка → тип использования, тип
            # продажи → тип контента, страна → территория КОДОМ («Germany» →
            # «DE», см. app/countries.py).
            "usage_type": {"column": 'Platform'},
            "content_type": {"column": 'Sales Type'},
            "territory": {"column": 'Country / Region', "as": "country"},
            # ВАЛЮТА: RU приходит в рублях, KZ в евро, AE в долларах. По этой
            # колонке сервис понимает, что нужен курс, и без него отчёт в
            # валюте не загрузит.
            "currency": {"column": 'Client Payment Currency'},
            # ПЕРИОД — ИЗ КОЛОНКИ «месяц отчёта»: в шапке файла его нет.
            "period": {"column": 'Reporting month'},
        },
    },
    {
        "name": 'Believe (русская шапка)',
        "signature": ('Каталожный номер релиза', 'Валюта', 'Сумма вознаграждения', 'Тип продажи', 'Платформа'),
        # Площадка в справочнике одна на все три отчёта (RU, KZ, AE): платежи
        # от Believe приходят от «BELEIVE DIGITAL» (опечатка в справочнике
        # живёт давно, и на неё завязано сопоставление платежей).
        "partner_names": ("BELEIVE DIGITAL",),
        "vat_rate": None,
        # Вид использования у Believe один на все строки — «streaming», как в
        # детализированном отчёте правообладателю из Dista. Остальные три
        # параметра идут ИЗ КОЛОНОК: у Believe в одном отчёте сотни сочетаний
        # площадки, типа продажи и страны.
        "attributes": {
            "usage_kind": "streaming",
        },
        "mapping": {
            # АРТИКУЛ — КАТАЛОЖНЫЙ НОМЕР РЕЛИЗА, и он главный (решение
            # владельца 24.09.2026: «именно по нашему артикулу мы сверяем в
            # первую очередь»). У альбома, сборника и YouTube-канала это номер
            # всей позиции (4100008 — канал Александра Иванова, 200 с лишним
            # видео), и деньги ложатся на неё, как в Dista. ISRC — запасной
            # путь, когда номера релиза нет в каталоге или он пуст.
            "sku": {"column": 'Каталожный номер релиза'},
            "code": {"column": 'ISRC'},
            "title": {"column": 'Название трека'},
            "artist": {"column": 'Исполнитель'},
            "quantity": {"column": 'Количество'},
            # ВСЁ — В СМЕЖНЫЕ: «Сумма вознаграждения» (Net Revenue) по формуле
            # Dista владельца. Механика (Mechanical Fee) не берётся: она почти
            # всегда ноль (39 строк на 534 ₽ в RU за май).
            "amount_related": {"column": 'Сумма вознаграждения'},
            # Параметры — по формуле Dista: площадка → тип использования, тип
            # продажи → тип контента, страна → территория КОДОМ («Germany» →
            # «DE», см. app/countries.py).
            "usage_type": {"column": 'Платформа'},
            "content_type": {"column": 'Тип продажи'},
            "territory": {"column": 'страна / регион', "as": "country"},
            # ВАЛЮТА: RU приходит в рублях, KZ в евро, AE в долларах. По этой
            # колонке сервис понимает, что нужен курс, и без него отчёт в
            # валюте не загрузит.
            "currency": {"column": 'Валюта'},
            # ПЕРИОД — ИЗ КОЛОНКИ «месяц отчёта»: в шапке файла его нет.
            "period": {"column": 'Месяц отчета'},
        },
    },
)


def match_builtin(columns: list) -> dict | None:
    """
    Готовое правило для этого файла — или None.

    ПРАВИЛА ПЛОЩАДОК ЖИВУТ В КОДЕ, а не заводятся руками на каждом партнёре
    (просьба владельца 18.09.2026: «на основе образца писать правила для всех
    отчётов»). Причины две. Формулу вроде «сумма × ставка авторских / сумма
    ставок» человек набирает, не видя данных, — и ошибка в ней выглядит как
    обычные числа, а не как поломка. И правило, единожды записанное в базу,
    остаётся там навсегда: поправишь его в коде — у партнёра всё равно
    сработает старая копия.

    Узнаётся файл ПО КОЛОНКАМ-ПРИМЕТАМ. Правило партнёра, если его завели
    руками, всё равно главнее: встроенное — это заготовка, а не запрет.
    """
    # КОЛОНКА-ПРИМЕТА ОБЯЗАНА БЫТЬ В ШАПКЕ ЕДИНСТВЕННОЙ (правка 24.09.2026,
    # найдено опытом на перекроенном отчёте МТС). Появись у площадки вторая
    # колонка с тем же именем ВЫШЕ настоящей — правило молча читало бы её:
    # артикулы в отчёте оказались бы мусором, а суммы при этом сошлись бы, и
    # заметить подмену было бы нечем. Пусть уж лучше файл не узнается вовсе:
    # тогда человеку покажут настройку колонок, и он увидит обе.
    #
    # Сюда приходят имена уже РАЗВЕДЁННЫЕ (`header_names`), поэтому повтор
    # виден как соседнее «X (2)»: есть оно — значит, имя «X» в этой шапке
    # адресом уже не является.
    keys = {normalize_header(c) for c in columns if _clean(c)}
    for rule in BUILTIN_RULES:
        if all(
            normalize_header(c) in keys
            and normalize_header(f"{c} (2)") not in keys
            for c in rule["signature"]
        ):
            # ПРИМЕТЫ БЫВАЮТ СЛИШКОМ ОБЩИМИ: «Код», «Количество», «Сумма»
            # найдутся в любой самодельной таблице. Такое правило помечено
            # `exact` и узнаёт файл, только если в шапке НИЧЕГО, кроме примет,
            # нет — иначе оно перехватывало бы чужие файлы.
            if rule.get("exact") and keys != {normalize_header(c) for c in rule["signature"]}:
                continue
            return rule
    return None


# ------------------------------------------------- период из шапки отчёта

# Основы названий месяцев: в отчётах они стоят и в именительном («июль 2026»),
# и в родительном («с 1 июля 2026»). Сравнивать целиком нельзя, а по основе —
# можно: «июн», «июл», «август»…
_MONTH_STEMS = (
    ("январ", 1), ("феврал", 2), ("март", 3), ("апрел", 4), ("ма", 5),
    ("июн", 6), ("июл", 7), ("август", 8), ("сентябр", 9), ("октябр", 10),
    ("ноябр", 11), ("декабр", 12),
)
_MONTH_RE = (
    "январ\\w*|феврал\\w*|март\\w*|апрел\\w*|ма[йяе]|июн\\w*|июл\\w*|"
    "август\\w*|сентябр\\w*|октябр\\w*|ноябр\\w*|декабр\\w*"
)
_ROMAN_QUARTERS = {"i": 1, "ii": 2, "iii": 3, "iv": 4}


def _month_number(word: str) -> int | None:
    """«июля» → 7. Сравниваем по основе: падеж у месяца в отчётах любой."""
    lowered = str(word or "").lower()
    for stem, number in _MONTH_STEMS:
        if lowered.startswith(stem):
            return number
    return None


def _month_end(year: int, month: int) -> date:
    return date(year + (month == 12), month % 12 + 1, 1) - timedelta(days=1)


def column_values(head: list, header_row: int, column: str) -> list:
    """
    Значения колонки в строках данных, попавших в верх файла (`head`).

    Для того, что одинаково во всём отчёте и видно по первым строкам: валюта,
    месяц отчёта. Весь файл ради этого читать незачем — у Believe он на
    полмиллиона строк.
    """
    if header_row >= len(head):
        return []
    key = normalize_header(column)
    names = [normalize_header(c) for c in header_names(head[header_row])]
    if key not in names:
        return []
    i = names.index(key)
    return [r[i] for r in head[header_row + 1:] if i < len(r) and _clean(r[i])]


_ISO_DAY = re.compile(r"(\d{4})[/.-](\d{1,2})[/.-](\d{1,2})")


def period_from_column(head: list, header_row: int, column: str) -> tuple | None:
    """
    Период отчёта ПО КОЛОНКЕ МЕСЯЦА — месяц целиком, от первого до последнего
    числа. У Believe периода в шапке нет вовсе, зато в каждой строке стоит
    «Месяц отчёта» («2026/05/01»): это и есть отчётный месяц (24.09.2026).
    Строки с разными месяцами дают период с первого по последний.
    """
    import calendar

    months = set()
    for value in column_values(head, header_row, column):
        if isinstance(value, datetime):
            months.add((value.year, value.month))
        elif isinstance(value, date):
            months.add((value.year, value.month))
        else:
            m = _ISO_DAY.search(_clean(value))
            if m and 1 <= int(m.group(2)) <= 12:
                months.add((int(m.group(1)), int(m.group(2))))
    if not months:
        return None
    first, last = min(months), max(months)
    return (
        date(first[0], first[1], 1),
        date(last[0], last[1], calendar.monthrange(last[0], last[1])[1]),
    )


def find_period(table: list, header_row: int) -> tuple | None:
    """
    Период отчёта, написанный В САМОМ ФАЙЛЕ, — пара дат или None.

    У МТС это строка над шапкой: «за период с 1 июля 2026 по 31 июля 2026».
    Читать её стоит потому, что период — единственное, что человек вводит
    руками на загрузке, и ошибиться в нём легче всего: файл за июнь грузят в
    июле, и месяц ставится «по умолчанию» не тот. В файле же он написан
    прямо.

    ЭТО ПОДСКАЗКА, А НЕ ПРИГОВОР: форма подставляет найденное, человек видит
    и может поправить. Поэтому и разбираем только очевидные записи, не пытаясь
    угадывать по обрывкам.
    """
    chunks = []
    for raw in table[:max(header_row, MAX_HEADER_SCAN)]:
        for cell in raw:
            value = _clean(cell)
            if value:
                chunks.append(value)
    if not chunks:
        return None
    text = " ".join(chunks).lower().replace("\xa0", " ")

    # «с 1 июля 2026 по 31 июля 2026» и «c 01 мая по 31 мая 2026г.»
    #
    # ГОД У ПЕРВОЙ ДАТЫ НЕОБЯЗАТЕЛЕН: у МегаФона он написан один раз, в
    # конце («период c 01 мая по 31 мая 2026г.»), и требовать его дважды
    # значит не прочитать период вовсе. Нет — берём год второй даты: период
    # внутри одного отчёта через новый год не переходит.
    #
    # «С» ЛОВИМ И КИРИЛЛИЦЕЙ, И ЛАТИНИЦЕЙ: в шапке МегаФона стоит латинская
    # «c» (U+0063). На вид не отличить, а регулярное выражение промахивается
    # молча — и период тихо не находится.
    #
    # ДЕНЬ БЫВАЕТ В КАВЫЧКАХ (ADV, 25.09.2026): «за период с "01" мая 2026 по
    # "31" мая 2026 г.». Без них выражение не срабатывало, и поиск падал до
    # «месяц ГОДА» — а тот находил строку выше, «к соглашению … от «01»
    # января 2019 года», то есть дату ДОГОВОРА вместо периода.
    q = r"[\"«»“”„']?"
    match = re.search(
        r"[сc]\s+" + q + r"(\d{1,2})" + q + r"\s+(" + _MONTH_RE + r")(?:\s+(\d{4}))?"
        r"\s+по\s+" + q + r"(\d{1,2})" + q + r"\s+(" + _MONTH_RE + r")\s+(\d{4})",
        text,
    )
    if match:
        d1, m1, y1, d2, m2, y2 = match.groups()
        first, second = _month_number(m1), _month_number(m2)
        if first and second:
            try:
                return date(int(y1 or y2), first, int(d1)), date(int(y2), second, int(d2))
            except ValueError:
                return None

    # «с 01.07.2026 по 31.07.2026», «01.07.2026 — 31.07.2026»
    match = re.search(
        r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})\s*(?:по|—|–|-)\s*"
        r"(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})",
        text,
    )
    if match:
        d1, m1, y1, d2, m2, y2 = (int(g) for g in match.groups())
        try:
            return date(y1, m1, d1), date(y2, m2, d2)
        except ValueError:
            return None

    # «ПО СОСТОЯНИЮ НА 01.08.2026» — ВОИС (правило владельца 25.09.2026):
    # отчёт за ПРОШЛЫЙ месяц, то есть весь июль. Месяц берём предыдущий
    # относительно даты, на которую составлен отчёт.
    match = re.search(
        r"по\s+состоянию\s+на\s+(\d{1,2})[.\-/](\d{1,2})[.\-/](\d{4})", text
    )
    if match:
        month, year = int(match.group(2)), int(match.group(3))
        if 1 <= month <= 12:
            month, year = (month - 1, year) if month > 1 else (12, year - 1)
            return date(year, month, 1), _month_end(year, month)

    # «III квартал 2026», «за 3 квартал 2026»
    match = re.search(r"\b(i{1,3}v?|[1-4])\s*квартал\w*\s*(\d{4})", text)
    if match:
        raw_quarter, year = match.groups()
        quarter = _ROMAN_QUARTERS.get(raw_quarter) or int(raw_quarter)
        start_month = (quarter - 1) * 3 + 1
        return date(int(year), start_month, 1), _month_end(int(year), start_month + 2)

    # «за июль 2026» — весь месяц
    match = re.search(r"\b(" + _MONTH_RE + r")\s+(\d{4})\b", text)
    if match:
        month = _month_number(match.group(1))
        if month:
            year = int(match.group(2))
            return date(year, month, 1), _month_end(year, month)
    return None


def suggest_mapping(columns: list) -> dict:
    """
    Догадка о правиле по названиям колонок — чтобы человеку не заполнять
    четыре поля с нуля на каждом новом партнёре.

    Это ПОДСКАЗКА: правило всё равно подтверждает человек, глядя на
    предпросмотр. Поэтому список синонимов короткий и без фантазии — то, что
    реально встречается в отчётах площадок.
    """
    # Слова взяты из настоящих отчётов, а не придуманы: у МТС артикул зовётся
    # «Код объекта Контента», количество — «Кол-во Продаж», название — «Название
    # объекта Контента». «№» стоит ПОСЛЕДНИМ: в отчёте МТС это порядковый номер
    # строки, а не артикул, и попасться на него легко.
    hints = {
        "sku": ("код объекта", "артикул", "код товара", "sku", "код", "номер", "№"),
        # ISRC и UPC — код САМОЙ ПЛОЩАДКИ, а не наш артикул: по нему трек
        # находят в каталоге, но в отчёт правообладателю уходит наш.
        "code": ("isrc", "upc"),
        "title": ("название объекта", "наименование", "название", "трек", "title", "track"),
        "artist": ("исполнитель", "артист", "artist", "performer"),
        "quantity": ("кол-во продаж", "количество", "кол-во", "прослушивания", "quantity", "streams"),
        "amount_author": ("сумма авт", "авторские", "сумма авторских", "author"),
        "amount_related": ("сумма смж", "смежные", "сумма смежных", "related", "master"),
    }
    mapping = {}
    used = set()
    # Идём ПО СЛОВАМ, а не по колонкам: порядок слов — это приоритет. Иначе
    # побеждает первый столбец файла, и в отчёте МТС артикулом становится «№»
    # (порядковый номер строки), потому что он стоит слева от «Кода объекта».
    for field_name, words in hints.items():
        for word in words:
            column = next(
                (
                    c
                    for c in columns
                    if c not in used
                    and normalize_header(c)
                    and (normalize_header(c) == word or normalize_header(c).startswith(word))
                ),
                None,
            )
            if column:
                mapping[field_name] = {"column": column}
                used.add(column)
                break
    return mapping
